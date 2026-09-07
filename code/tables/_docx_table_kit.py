# -*- coding: utf-8 -*-
"""Shared python-docx helpers for the submission tables.

Used by 08_22_2026_table1_word.py and 08_22_2026_table2_word.py.  These are the bits that
need raw OOXML poking, which python-docx does not expose - factored out so the two tables
cannot drift apart on the details that actually decide whether a table survives pagination.

⚠ ARIAL HAS NO U+21B3 (the ↳ used for conditional items in the HTML artifacts).  It prints
as a missing-glyph box.  Show conditional rows by INDENTATION instead.  `assert_glyphs()`
below makes that failure loud rather than visual: pass it every non-ASCII character you are
about to typeset.
"""
import os
from docx.shared import Pt, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TWIP = 635                                   # EMU per twip
dxa = lambda n: Emu(int(n * TWIP))

INK = RGBColor(0x1B, 0x20, 0x29)
MUT = RGBColor(0x4A, 0x52, 0x59)
FAINT = RGBColor(0x8B, 0x95, 0xA1)
ACCENT = RGBColor(0x16, 0x69, 0x7A)
PROB = RGBColor(0xA6, 0x43, 0x2A)            # the artifacts' --prob, for the problem response
BAND = "F2F5F6"

LETTER_W, LETTER_H, MARGIN = 12240, 15840, 864


def assert_glyphs(text, font=r"C:\Windows\Fonts\arial.ttf"):
    """Fail loudly on a character the output font cannot draw.  Arial's missing ↳ reached a
    built figure once already; a box in a PDF is not something a build otherwise notices."""
    try:
        from fontTools.ttLib import TTFont
    except ImportError:
        return                                # not installed: skip rather than block the build
    cmap = set(TTFont(font).getBestCmap())
    missing = sorted({c for c in text if ord(c) > 127 and ord(c) not in cmap})
    assert not missing, ("font cannot draw: " + " ".join("U+%04X %s" % (ord(c), c) for c in missing))


def new_document(landscape=False, margin=MARGIN, font="Arial", size=8):
    from docx import Document
    doc = Document()
    sec = doc.sections[0]
    if landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = dxa(LETTER_H), dxa(LETTER_W)
    else:
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = dxa(LETTER_W), dxa(LETTER_H)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, dxa(margin))
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = font, Pt(size)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    rpr = normal.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(a), font)
    usable = (sec.page_width.emu // TWIP) - 2 * margin
    return doc, usable


def para(cellobj, runs, align=WD_ALIGN_PARAGRAPH.RIGHT, indent=0, space_after=0, font="Arial"):
    """runs = [(text, size_pt, colour, bold)]"""
    p = cellobj.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.left_indent = dxa(indent)
    for txt, size, col, bold in runs:
        r = p.add_run(txt)
        r.font.name, r.font.size, r.font.color.rgb, r.bold = font, Pt(size), col, bold
    return p


def shade(cellobj, hexfill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")              # never SOLID - it renders black
    el.set(qn("w:fill"), hexfill)
    cellobj._tc.get_or_add_tcPr().append(el)


def borders(cellobj, **edges):
    """edges: top/bottom/left/right -> (size_in_eighths_of_a_point, hexcolour) or None for nil"""
    tcpr = cellobj._tc.get_or_add_tcPr()
    tb = tcpr.find(qn("w:tcBorders"))
    if tb is None:
        tb = OxmlElement("w:tcBorders"); tcpr.append(tb)
    for edge, spec in edges.items():
        e = OxmlElement("w:" + edge)
        if spec is None:
            e.set(qn("w:val"), "nil")
        else:
            sz, col = spec
            e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:color"), col)
        tb.append(e)


def repeat_header(row):
    """Repeat this row at the top of every page the table spans."""
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trpr.append(el)


def cant_split(row):
    """A row must not break across pages - half a row of digits is unreadable."""
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); trpr.append(el)


def keep_next(row):
    """Bind a row to the one after it.

    ⚠ This is what stops a heading - or a repeated header block - being stranded at a page
    foot.  Table 1's first build split a panel BETWEEN its two header rows, which also
    defeated repeat_header(): the repeated block was itself the thing that broke."""
    for c in row.cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = True


def render_and_count_pages(docx_path, pdf_path):
    """Convert through Word and return the page count.  There is no LibreOffice on this
    machine; Word is driven over COM from PowerShell.  Returns None if it is unavailable -
    the build should not fail because a verification convenience is missing."""
    import subprocess, tempfile
    ps = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; $w.DisplayAlerts = 0; "
        "$d = $w.Documents.Open('%s', $false, $true); $d.SaveAs([ref]'%s', [ref]17); "
        "$d.ComputeStatistics(2); $d.Close($false); $w.Quit()"
        % (os.path.abspath(docx_path).replace("'", "''"), os.path.abspath(pdf_path).replace("'", "''"))
    )
    try:
        out = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", ps],
                             capture_output=True, text=True, timeout=180)
        return int(out.stdout.strip().splitlines()[-1])
    except Exception:
        return None
