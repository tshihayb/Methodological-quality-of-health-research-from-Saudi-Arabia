# NOTE (public repository): 14 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
# -*- coding: utf-8 -*-
"""Build the corr-author -> Non-Saudi_corr re-check trace (per-paper) + summary workbook."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

def G(r,c):
    v=str(r[c]).strip() if c in r.index else ''
    return '' if v.lower()=='nan' else v
def is_corr(s): return 'corresponding author' in s.lower() or 'not from a saudi' in s.lower()

P385=set(pd.read_csv('data/analysis/07_23_2026_ANALYSIS_DATASET_wide.csv',dtype=str,keep_default_na=False).PMID.astype(str))
TARGET_EXCL={'STUDY-0439','STUDY-0274','STUDY-0020','STUDY-0082','STUDY-0480','STUDY-0187','STUDY-0033',
             'STUDY-0509','STUDY-0316','STUDY-0448','STUDY-0593','STUDY-0785','STUDY-0097'}
COI={'STUDY-0956'}

exc=pd.read_excel('data/screening/f300_ta_ya_TA.xlsx','Excluded',dtype=str).fillna('')
tcorr=pd.read_excel('data/screening/f300_ta_ya_TA.xlsx','T_corr',dtype=str).fillna('')
tcorr_map={G(r,'PMID'):(G(r,'Status_t_new'),G(r,'Reason for Exclusion_t_new') or G(r,'Study Type_t_new')) for _,r in tcorr.iterrows()}
ns=pd.read_excel('data/screening/f300_ta_ya_clean.xlsx','Non-Saudi_corr',dtype=str).fillna('')
ns_map={G(r,'PMID'):dict(sy=G(r,'Status_y'),st=G(r,'Status_t'),tyy=G(r,'Study Type_y_new'),
                         tyt=G(r,'Study Type_t_new'),agr=G(r,'Agreement'),adj=G(r,'Adjudicaiton')) for _,r in ns.iterrows()}
ns_ids=set(ns_map)

rows=[]
for _,r in exc.iterrows():
    pm=G(r,'PMID'); rt=G(r,'Reason for Exclusion_t'); ry=G(r,'Reason for Exclusion_y')
    if not (is_corr(rt) or is_corr(ry)): continue
    reached = pm in ns_ids
    # why held back
    if reached:
        held=''
    elif is_corr(rt) and is_corr(ry):
        tc=tcorr_map.get(pm,('','')); held=f'both wrote corr-author; T_corr re-adjudicated -> {tc[0]}/{tc[1]}'
    else:
        other = ry if (is_corr(rt) and not is_corr(ry)) else rt
        held=f'other reviewer already had a real reason: {other}'
    # pool decision
    p=ns_map.get(pm)
    if p:
        adj=p['adj']
        if adj.lower().startswith('excl'): pool_dec='EXCLUDE'
        elif p['sy']=='Include' and p['st']=='Include': pool_dec='Include'
        elif adj: pool_dec='Include ('+adj+')'
        else: pool_dec=f"{p['sy']}/{p['st']}"
        pool_reason = (p['tyy'] if p['sy']=='Exclude' else (p['tyt'] if p['st']=='Exclude' else ''))
    else:
        pool_dec=''; pool_reason=''
    dispo = ('INCLUDED (385)' if pm in P385 else 'COI' if pm in COI else
             'target-reached' if pm in TARGET_EXCL else 'EXCLUDED (screening)')
    rows.append(dict(PMID=pm,link=f'https://pubmed.ncbi.nlm.nih.gov/{pm}/',TA_reason=rt,YA_reason=ry,
        reached_recheck_pool=('Yes' if reached else 'No'), held_back_because=held,
        pool_decision=pool_dec, pool_recorded_reason=pool_reason,
        final_disposition=dispo, in_385=('Yes' if pm in P385 else '')))
tr=pd.DataFrame(rows).sort_values('reached_recheck_pool')
print("corr-author exclusions traced:",len(tr))
print("reached pool:", (tr.reached_recheck_pool=='Yes').sum(), "| held back:", (tr.reached_recheck_pool=='No').sum())
print("\npool decisions:"); print(tr[tr.reached_recheck_pool=='Yes'].pool_decision.map(lambda s:'EXCLUDE' if s=='EXCLUDE' else 'Include').value_counts().to_string())
print("\nof reached-pool, final disposition:"); print(tr[tr.reached_recheck_pool=='Yes'].final_disposition.value_counts().to_string())
print("\nof held-back, final disposition:"); print(tr[tr.reached_recheck_pool=='No'].final_disposition.value_counts().to_string())
print("\nheld-back reasons:");
for x in tr[tr.reached_recheck_pool=='No'].held_back_because: print('   -',x)

# summary
summ=pd.DataFrame({'Metric':[
    'Corr-author exclusions (union of TA + YA "corr-author" calls)',
    '  reached the Non-Saudi_corr re-check pool',
    '  held back (another exclusion criterion already applied)',
    '     - held back: other reviewer gave a substantive reason',
    '     - held back: both wrote corr-author, T_corr re-adjudicated to a real reason',
    'Re-check pool outcomes (of the 70)',
    '  re-Included -> now in the 385',
    '  re-Included then removed for COI (STUDY-0956)',
    '  adjudicated back to Exclude (non-human x3, "narrative review"/protocol x1)'],
    'n':[len(tr),(tr.reached_recheck_pool=='Yes').sum(),(tr.reached_recheck_pool=='No').sum(),
         sum('other reviewer' in x for x in tr.held_back_because),
         sum('both wrote' in x for x in tr.held_back_because),70,
         ((tr.reached_recheck_pool=='Yes')&(tr.in_385=='Yes')).sum(),
         ((tr.reached_recheck_pool=='Yes')&(tr.final_disposition=='COI')).sum(),
         ((tr.reached_recheck_pool=='Yes')&(tr.pool_decision=='EXCLUDE')).sum()]})

out='data/screening/08_13_2026_corr_author_recheck_trace.xlsx'
with pd.ExcelWriter(out,engine='openpyxl') as xw:
    summ.to_excel(xw,sheet_name='Summary',index=False)
    tr.to_excel(xw,sheet_name='Trace_101',index=False)
    hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='16697A')
    for sh in xw.sheets:
        ws=xw.sheets[sh]; ws.freeze_panes='A2'
        for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(vertical='top',wrap_text=True)
    for shn,df in [('Summary',summ),('Trace_101',tr)]:
        ws=xw.sheets[shn]
        for i,col in enumerate(df.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(60,max(10,int(df[col].map(lambda v:len(str(v))).max())+2))
print('\nwrote',out)
