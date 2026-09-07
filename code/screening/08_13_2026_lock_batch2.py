# NOTE (public repository): 26 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
# -*- coding: utf-8 -*-
"""Lock batch-2 exclusion reasons: apply decisions, reclassify the 2 exp/outcome 'full-text' papers
to Non-health, tag the 5 exp/outcome exclusions, refresh audit + master CSV."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NH='Non-human / laboratory'; REV='Review (narrative/systematic)'; QUAL='Qualitative research'
NE='Non-empirical (protocol/simulation)'; CR='Case report / series'; NHL='Non-health topic'

rec = pd.read_csv('data/screening/08_13_2026_exclusion_reason_recount_audit.csv', dtype=str).fillna('')
b2 = rec[rec.batch=='batch-2'].copy()
b2['Seq_n']=pd.to_numeric(b2['Seq'],errors='coerce')

B2 = {
 'STUDY-0582':(NH ,'disagreement -> TA Non-human (computational SARS-CoV-2 drug study; biomedical in-silico)'),
 'STUDY-0684':(NE ,'disagreement -> TA Non-empirical (forensic-training commentary; not original research)'),
 'STUDY-0858':(NHL,'harmonize Non-health (glossary of open-scholarship terms; not health/not a study)'),
 'STUDY-0998':(NE ,'reviewer comment "RCT protocol" -> Non-empirical, not Qualitative'),
 'STUDY-0545':(NHL,'reviewer comment "cyber security...not health" -> Non-health'),
 'STUDY-0873':(NHL,'reviewer comment "languages research" -> Non-health'),
 'STUDY-0754':(NHL,'reviewer comment "not health research" -> Non-health'),
 'STUDY-0483':(NHL,'reviewer comment "non-health" -> Non-health'),
 'STUDY-0323':(NH ,'TA Non-human (YA blank) [FLAG: "used in reviewer testing round 1" — confirm in denominator]'),
 'STUDY-0874':(NH ,'TA Non-human (YA blank)'),
 'STUDY-0726':(QUAL,'TA Qualitative (YA blank)'),
 'STUDY-0470':(NH ,'TA Non-human (YA blank)'),
 'STUDY-0054':(NH ,'TA Non-human (YA blank)'),
 'STUDY-0435':(NHL,'TA Non-health (YA blank)'),
 'STUDY-0652':(NH ,'TA Non-human (YA blank)'),
 'STUDY-0278':(NHL,'harmonize Non-health (Sensors: electricity load forecasting)'),
 'STUDY-0602':(NHL,'harmonize Non-health (Sensors: Cloud-of-Things crowd monitoring)'),
 'STUDY-0353':(NHL,'harmonize Non-health (blockchain economics)'),
 'STUDY-0534':(NHL,'harmonize Non-health (ML for environmental pollution)'),
 'STUDY-0392':(NHL,'harmonize Non-health (Sensors: LiDAR assistive edge device; engineering) [BORDERLINE-user OK]'),
}
b2['FINAL_reason']=b2.apply(lambda r: B2[r['PMID']][0] if r['PMID'] in B2 else r['RESOLVED_reason'], axis=1)
b2['basis']=b2.apply(lambda r: B2[r['PMID']][1] if r['PMID'] in B2 else 'agreed / adjudicated on record', axis=1)

# reason-source label
def has(s): return str(s).strip()!='' and str(s).strip().lower()!='nan'
def rsource(r):
    ta,ya=has(r['TA_reason_orig']),has(r['YA_reason_orig'])
    if ta and ya: return 'both reviewers'
    if ta: return 'TA only'
    if ya: return 'YA only'
    return 'adjudicated (neither wrote a reason)'
b2['reason_source']=b2.apply(rsource,axis=1)

# ---- add the 2 exp/outcome papers reclassified from 'full-text-not-found' -> Non-health ----
addrows=[]
for pm,seq,ty,radj in [('STUDY-0500','324','Causal','Non-health related'),
                       ('STUDY-0080','441','Descriptive','Non-health (health economic; hard to ascertain exp/out)')]:
    addrows.append({'Seq':seq,'PMID':pm,'batch':'batch-2','TA_reason_orig':'','YA_reason_orig':'',
        'RESOLVED_reason':NHL,'flag_tier':'','FINAL_reason':NHL,'Seq_n':float(seq),
        'basis':f'excluded during exp/outcome determination -> {radj} (was mislabelled full-text-not-found)',
        'reason_source':'adjudicated (exp/outcome determination)'})
b2 = pd.concat([b2, pd.DataFrame(addrows)], ignore_index=True).sort_values('Seq_n')

# tag the 5 exp/outcome (extraction-step) exclusions
EXPOUT={'STUDY-0500','STUDY-0080','STUDY-0846','STUDY-0586','STUDY-0668'}
b2['extraction_step']=b2.PMID.map(lambda p:'excluded at data-extraction (exp/outcome)' if p in EXPOUT else '')

ORDER=[NH,NHL,REV,QUAL,NE,CR]
t=b2['FINAL_reason'].value_counts()
print("="*56); print("BATCH-2 LOCKED — %d screening/eligibility exclusions"%len(b2)); print("="*56)
tot=0
for k in ORDER:
    print(f"  {k:<38} {int(t.get(k,0)):>4}"); tot+=int(t.get(k,0))
print(f"  {'TOTAL':<38} {tot:>4}")
assert tot==128, tot
print("\nreason_source:"); print(b2.reason_source.value_counts().to_string())

# ---- workbook ----
b2out=b2[['Seq','PMID','batch','TA_reason_orig','YA_reason_orig','reason_source','FINAL_reason','extraction_step','basis']].copy()
b2out['link']='https://pubmed.ncbi.nlm.nih.gov/'+b2out['PMID']+'/'
b2out=b2out[['Seq','PMID','link','TA_reason_orig','YA_reason_orig','reason_source','FINAL_reason','extraction_step','basis']]
summ=pd.DataFrame([{'Exclusion reason':k,'n':int(t.get(k,0))} for k in ORDER]+[{'Exclusion reason':'TOTAL','n':128}])
src=b2.reason_source.value_counts().reset_index(); src.columns=['reason_source','n']
readme=pd.DataFrame({'Item':['BATCH-2 EXCLUSIONS — LOCKED','Prepared','Scope','Decisions','','','','Borderlines (my calls)','','Note'],
 'Detail':['Final reason for every batch-2 (Seq 301-621) screening/eligibility exclusion.','2026-08-13',
   'Batch 2 = 321 screened -> 191 included, 130 excluded (128 with reasons + 2 no-Saudi-affiliation).',
   '(1) reason-disagreements -> take one reviewer reason; (2) where a reviewer wrote a comment, followed the comment; '
   '(3) harmonized Non-health for clearly non-health papers (Sensors/IoT/blockchain/environmental); '
   '(4) YA-blank tail -> TA reason.',
   '(5) STUDY-0500 + STUDY-0080: both Include/Include then excluded at exposure/outcome determination for Non-health '
   '-> reclassified from "full-text-not-found" to Non-health (consistent with the other 3 exp/outcome exclusions).',
   'The 5 exp/outcome papers (col extraction_step) were included at screening and excluded later during data extraction; '
   'they can be shown as a distinct "excluded at data extraction (n=5)" step in the flow if preferred.','',
   'STUDY-0392 -> Non-health (LiDAR assistive device); STUDY-0635 kept Non-human (biosensor material); '
   'STUDY-0323 kept Non-human but flagged "used in reviewer testing round 1" (confirm whether it belongs in the denominator).','',
   'reason_source: 101 both-agreed, 3 both-disagreed, 20 single-reviewer, 2+2 adjudicated (incl. the exp/outcome reclassifications).']})

out='data/screening/08_13_2026_batch2_exclusions_LOCKED.xlsx'
with pd.ExcelWriter(out,engine='openpyxl') as xw:
    readme.to_excel(xw,sheet_name='README',index=False)
    summ.to_excel(xw,sheet_name='Tally',index=False)
    src.to_excel(xw,sheet_name='Reason_source',index=False)
    b2out.to_excel(xw,sheet_name='Batch2_audit_128',index=False)
    hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='16697A')
    for sh in xw.sheets:
        ws=xw.sheets[sh]; ws.freeze_panes='A2'
        for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(vertical='top',wrap_text=True)
    xw.sheets['README'].column_dimensions['A'].width=24; xw.sheets['README'].column_dimensions['B'].width=120
    for c in xw.sheets['README']['A']: c.font=Font(bold=True); c.alignment=Alignment(vertical='top')
    for c in xw.sheets['README']['B']: c.alignment=Alignment(vertical='top',wrap_text=True)
    for shn,df in [('Tally',summ),('Reason_source',src),('Batch2_audit_128',b2out)]:
        ws=xw.sheets[shn]
        for i,col in enumerate(df.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(70,max(10,int(df[col].map(lambda v:len(str(v))).max())+2))
print("\nwrote",out)

# ---- update master CSV (batch-1 already LOCKED there; refresh batch-2) ----
m=pd.read_csv('data/screening/08_13_2026_exclusion_final_master.csv',dtype=str).fillna('')
m=m[m.batch!='batch-2']              # drop old provisional batch-2 rows
keep=['Seq','PMID','batch','TA_reason_orig','YA_reason_orig','RESOLVED_reason','flag_tier','FINAL_reason']
b2m=b2.copy(); b2m['status']='batch-2 LOCKED'
for c in keep:
    if c not in b2m.columns: b2m[c]=''
master=pd.concat([m, b2m[keep+['status']]], ignore_index=True)
master.to_csv('data/screening/08_13_2026_exclusion_final_master.csv',index=False,encoding='utf-8-sig')
print("updated data/screening/08_13_2026_exclusion_final_master.csv  (batch-1 + batch-2 both LOCKED)")
