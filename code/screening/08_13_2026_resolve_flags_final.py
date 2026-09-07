# NOTE (public repository): 24 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
# -*- coding: utf-8 -*-
"""Apply best-effort resolutions to the 24 flagged papers (from abstract review) and
produce the FINAL fully-resolved exclusion-reason breakdown + updated audit workbook."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

rec = pd.read_csv('data/screening/08_13_2026_exclusion_reason_recount_audit.csv', dtype=str).fillna('')
rec['Seq_n'] = pd.to_numeric(rec['Seq'], errors='coerce')
rec['link'] = 'https://pubmed.ncbi.nlm.nih.gov/' + rec['PMID'] + '/'

NH='Non-human / laboratory'; REV='Review (narrative/systematic)'; QUAL='Qualitative research'
NE='Non-empirical (protocol/simulation)'; CR='Case report / series'; NHL='Non-health topic'

# best-effort resolution for the 24 flagged, each with the deciding evidence
RESOLVE = {
 # --- Tier D (disposition uncertain -> resolved via abstract; all have a Saudi author) ---
 'STUDY-0374': (NH,  'In-silico TCGA gene-expression / PPI-network analysis (no human subjects)'),
 'STUDY-0168': (NH,  'In-vitro (J774.2 macrophages) + molecular docking of smokeless-tobacco samples'),
 'STUDY-0849': (NH,  'Optical modelling of blue-blocking lenses; no human participants'),
 'STUDY-0769': (NH,  'DeepSVP — deep-learning method paper (Bioinformatics), computational'),
 'STUDY-0683': (NE,  'HEROES "study protocol" — no results reported'),
 # --- Tier A (reason disagreement -> resolved via title/journal) ---
 'STUDY-0836': (NHL, 'Vehicular-network engineering (Sensors); not health research'),
 'STUDY-0190': (NH,  'In-silico pharmacophore drug screening vs HIV-1 protein'),
 'STUDY-0219': (NH,  'Environmental groundwater contamination sampling (health-risk assessment)'),
 'STUDY-0681': (NHL, 'Healthcare cybersecurity / IT-systems paper ([NAME-REDACTED]); not health research'),
 'STUDY-0582': (NH,  'Computational (docking) study of drugs vs SARS-CoV-2'),
 'STUDY-0684': (NE,  'Forensic-training commentary, 3 pp (Med Leg J); not original research'),
 'STUDY-0858': (NHL, 'Community-sourced glossary of open-scholarship terms; not health, not a study'),
 # --- Tier B (reason vs comment -> resolved) ---
 'STUDY-0998': (NE,  'Home-based PA programme in T2DM — study PROTOCOL (comment: "RCT protocol")'),
 'STUDY-0545': (NHL, 'Homomorphic-encryption patient-data privacy (crypto); comment "cyber security, not health"'),
 'STUDY-0873': (NHL, 'Sociolinguistics — attitudes toward Saudi English; comment "languages research"'),
 'STUDY-0754': (NHL, 'Smart-city IoT/IT (Sensors); comment "not health research"'),
 'STUDY-0483': (NHL, 'Face-recognition in video (Sensors); comment "non-health"'),
 # --- Tier C (TA-only, YA status blank -> use TA reason) ---
 'STUDY-0323': (NH,  'TA: Non-human (comment: paper used in reviewer testing round 1)'),
 'STUDY-0874': (NH,  'TA: Non-human (YA left status blank)'),
 'STUDY-0726': (QUAL,'TA: Qualitative (YA left status blank)'),
 'STUDY-0470': (NH,  'TA: Non-human (YA left status blank)'),
 'STUDY-0054': (NH,  'TA: Non-human (YA left status blank)'),
 'STUDY-0435': (NHL, 'TA: Non-health (YA left status blank)'),
 'STUDY-0652': (NH,  'TA: Non-human (YA left status blank)'),
}
def resolved_reason(r):
    if r['flag_tier'] == '':
        return r['FINAL_reason'], 'agreed / adjudicated on record'
    return RESOLVE.get(r['PMID'], (r['FINAL_reason'], 'UNRESOLVED'))
rr = rec.apply(resolved_reason, axis=1)
rec['RESOLVED_reason'] = [x[0] for x in rr]
rec['resolution_evidence'] = [x[1] for x in rr]

ORDER=[NH,REV,QUAL,NE,NHL,CR]
tally = rec['RESOLVED_reason'].value_counts()
print("="*60); print("FINAL fully-resolved breakdown (n=%d)"%len(rec)); print("="*60)
tot=0
for k in ORDER:
    print(f"  {k:<38} {int(tally.get(k,0)):>4}"); tot+=int(tally.get(k,0))
print(f"  {'TOTAL':<38} {tot:>4}")
assert tot==233, tot
# any leftover categories?
extra=[k for k in tally.index if k not in ORDER]
if extra: print("  UNEXPECTED:", {k:int(tally[k]) for k in extra})

# comparison table
prov={'Non-human / laboratory':178,REV:14,'Corr. author not Saudi':12,QUAL:11,NE:8,CR:6,NHL:1,'Reason not recorded':3}
print("\nPROVISIONAL -> CORRECTED")
for k in [NH,REV,'Corr. author not Saudi',QUAL,NE,CR,NHL,'Reason not recorded']:
    print(f"  {k:<34} {prov.get(k,0):>4}  ->  {int(tally.get(k,0)) if k in ORDER else 0:>4}")

# ---------------- rewrite workbook with resolved columns ----------------
AC='16697A'
clean_n=int((rec.flag_tier=='').sum()); flag_n=int((rec.flag_tier!='').sum())
summ=pd.DataFrame([{'Exclusion reason':k,'Provisional':prov.get(k,0),'Corrected (final)':int(tally.get(k,0))} for k in ORDER]
    +[{'Exclusion reason':'Corr. author not Saudi (REMOVED — criterion relaxed to any-Saudi-author)','Provisional':12,'Corrected (final)':0},
      {'Exclusion reason':'Reason not recorded (RESOLVED)','Provisional':3,'Corrected (final)':0},
      {'Exclusion reason':'TOTAL screening exclusions','Provisional':233,'Corrected (final)':233}])

readme=pd.DataFrame({'Item':['EXCLUSION-REASON RECOUNT (final)','Prepared','Rule','Scope','Ground truth',
    'Headline change','','On the 24 flagged','Confidence','','Columns'],
 'Detail':['Reason each screening-excluded paper was excluded, under the agreed/adjudicated rule.','2026-08-13',
   'Reason where both reviewers agreed; adjudicated reason where they disagreed. Batch-1 folds in the T_corr re-adjudication of the relaxed corresponding-author filter.',
   '233 screening exclusions (batch-1=107, batch-2=126). Separate PRISMA steps not counted here: full-text-not-found 2, target-reached 13, 3-more panel 4, no-Saudi-affiliation 2, COI 1.',
   'Canonical 385 analysis set decides inclusion.',
   'The first-pass "corresponding-author-not-Saudi" filter was relaxed to "any Saudi author"; those 78 papers were re-adjudicated (T_corr). NO standalone corr-author category remains (provisional had 12). Non-health rose 1->%d; non-human 178->%d; non-empirical 8->%d.'%(int(tally.get(NHL,0)),int(tally.get(NH,0)),int(tally.get(NE,0))),
   '',
   '%d assigned cleanly (agreed/adjudicated on record). %d needed best-effort coding from the PubMed title/abstract — see RESOLVED_reason + resolution_evidence; please verify these.'%(clean_n,flag_n),
   'The 24 best-effort codings are documented per-paper with the deciding evidence; totals shift by at most a few papers between adjacent categories if any are re-coded.',
   '',
   'flag_tier: blank=clean; A=reason disagreement; B=reason-vs-comment; C=TA-only(YA blank); D=disposition-uncertain(now resolved).']})

audit=rec.sort_values(['batch','Seq_n'])[['Seq','PMID','link','batch','TA_reason_orig','YA_reason_orig','tcorr',
    'assignment_method','flag_tier','RESOLVED_reason','resolution_evidence','evidence']]
flagged=rec[rec.flag_tier!=''].sort_values(['flag_tier','batch','Seq_n'])[['flag_tier','Seq','PMID','link','batch',
    'TA_reason_orig','YA_reason_orig','tcorr','RESOLVED_reason','resolution_evidence','evidence']]

out='data/screening/08_13_2026_exclusion_reason_audit.xlsx'
with pd.ExcelWriter(out, engine='openpyxl') as xw:
    readme.to_excel(xw,sheet_name='README',index=False)
    summ.to_excel(xw,sheet_name='Summary_tally',index=False)
    audit.to_excel(xw,sheet_name='Full_audit_233',index=False)
    flagged.to_excel(xw,sheet_name='Best_effort_24',index=False)
    hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor=AC)
    for sh in xw.sheets:
        ws=xw.sheets[sh]; ws.freeze_panes='A2'
        for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(vertical='top',wrap_text=True)
    xw.sheets['README'].column_dimensions['A'].width=24; xw.sheets['README'].column_dimensions['B'].width=120
    for c in xw.sheets['README']['A']: c.font=Font(bold=True); c.alignment=Alignment(vertical='top')
    for c in xw.sheets['README']['B']: c.alignment=Alignment(vertical='top',wrap_text=True)
    for shn,df in [('Summary_tally',summ),('Full_audit_233',audit),('Best_effort_24',flagged)]:
        ws=xw.sheets[shn]
        for i,col in enumerate(df.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(58,max(11,int(df[col].map(lambda v:len(str(v))).max())+2))
print('\nwrote', out)
rec.to_csv('data/screening/08_13_2026_exclusion_reason_recount_audit.csv', index=False, encoding='utf-8-sig')
