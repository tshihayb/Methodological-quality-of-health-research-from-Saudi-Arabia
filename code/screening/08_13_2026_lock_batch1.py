# NOTE (public repository): 17 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
# -*- coding: utf-8 -*-
"""Lock batch-1 exclusion reasons: apply all user decisions, fold in STUDY-0097,
apply harmonized Non-health category, refresh the batch-1 audit + master CSV."""
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NH='Non-human / laboratory'; REV='Review (narrative/systematic)'; QUAL='Qualitative research'
NE='Non-empirical (protocol/simulation)'; CR='Case report / series'; NHL='Non-health topic'

rec = pd.read_csv('data/screening/08_13_2026_exclusion_reason_recount_audit.csv', dtype=str).fillna('')
rec['Seq_n']=pd.to_numeric(rec['Seq'],errors='coerce')

# ---- batch-1 FINAL overrides (PMID -> (final_reason, basis)) ----
B1 = {
 'STUDY-0374':(NH ,'Non-Saudi_corr pool: adjudicated Exclude, YA Non-human (in-silico TCGA)'),
 'STUDY-0168':(NH ,'Non-Saudi_corr pool: adjudicated Exclude, YA Non-human (in-vitro + docking)'),
 'STUDY-0849':(NH ,'Non-Saudi_corr pool: adjudicated Exclude, YA Non-human (optical modelling)'),
 'STUDY-0683':(REV,'Non-Saudi_corr pool: adjudicated Exclude, YA "Narrative review" (factually a study protocol)'),
 'STUDY-0769':(NH ,'Saudi_corr: adjudicated Exclude, YA Non-human (DeepSVP deep-learning method)'),
 'STUDY-0190':(NH ,'reason-disagreement -> TA reason Non-human (in-silico HIV drug screening; biomedical)'),
 'STUDY-0219':(NH ,'reason-disagreement -> Non-human (Chemosphere groundwater + health-risk; user kept Non-human)'),
 'STUDY-0836':(NHL,'harmonized Non-health (Sensors: vehicular-network engineering)'),
 'STUDY-0681':(NHL,'harmonized Non-health ([NAME-REDACTED]: medical-data encryption / IT security)'),
 'STUDY-0404':(NHL,'harmonized Non-health (Sensors: Ant-Miner data-mining classifier)'),
 'STUDY-0881':(NHL,'harmonized Non-health (Sensors: wireless-sensor-network authentication)'),
 'STUDY-0768':(NHL,'harmonized Non-health ([NAME-REDACTED]: healthcare-data steganography)'),
 'STUDY-0363':(NHL,'harmonized Non-health (Comput Intell Neurosci: medical-signal cryptography)'),
 'STUDY-0724':(NHL,'harmonized Non-health (Chemosphere: Cd/Pb adsorbent materials for water)'),
 'STUDY-0301':(NHL,'harmonized Non-health (Chemosphere: phytoextraction / soil agronomy)'),
}

b1 = rec[rec.batch=='batch-1'].copy()
b1['FINAL_reason']=b1.apply(lambda r: B1[r['PMID']][0] if r['PMID'] in B1 else r['RESOLVED_reason'], axis=1)
b1['basis']=b1.apply(lambda r: B1[r['PMID']][1] if r['PMID'] in B1 else 'agreed / adjudicated on record', axis=1)

# ---- fold in STUDY-0097 (was mislabelled target-reached; screened at Seq 181, Non-human) ----
extra = pd.DataFrame([{'Seq':'181','PMID':'STUDY-0097','batch':'batch-1',
   'TA_reason_orig':'Non-human research','YA_reason_orig':'Corresponding author was not from a Saudi institution',
   'RESOLVED_reason':NH,'flag_tier':'','FINAL_reason':NH,
   'basis':'folded from target-reached: batch-1 Seq 181 + panel both Non-human (Cryptosporidium parasite genomics)',
   'Seq_n':181.0}])
b1 = pd.concat([b1, extra], ignore_index=True).sort_values('Seq_n')

ORDER=[NH,NHL,REV,QUAL,NE,CR]
t=b1['FINAL_reason'].value_counts()
print("="*56); print("BATCH-1 LOCKED — %d screening exclusions"%len(b1)); print("="*56)
tot=0
for k in ORDER:
    print(f"  {k:<38} {int(t.get(k,0)):>4}"); tot+=int(t.get(k,0))
print(f"  {'TOTAL':<38} {tot:>4}")
assert tot==108, tot
extra_cats=[k for k in t.index if k not in ORDER]
if extra_cats: print("  UNEXPECTED:",{k:int(t[k]) for k in extra_cats})

# ---- write batch-1 audit workbook ----
b1out=b1[['Seq','PMID','batch','TA_reason_orig','YA_reason_orig','FINAL_reason','basis']].copy()
b1out['link']='https://pubmed.ncbi.nlm.nih.gov/'+b1out['PMID']+'/'
b1out=b1out[['Seq','PMID','link','TA_reason_orig','YA_reason_orig','FINAL_reason','basis']]
summ=pd.DataFrame([{'Exclusion reason':k,'n':int(t.get(k,0))} for k in ORDER]+[{'Exclusion reason':'TOTAL','n':108}])
readme=pd.DataFrame({'Item':['BATCH-1 EXCLUSIONS — LOCKED','Prepared','Scope','Decisions applied','','','','','Note'],
 'Detail':['Final reason for every batch-1 (Seq 1-300) screening exclusion.','2026-08-13',
   '108 screening exclusions of the 300 batch-1 papers (191 included in the 385; 1 removed for COI = STUDY-0956).',
   '(1) corr-author re-check folded in via Non-Saudi_corr/Saudi_corr pools (authoritative joint adjudication);',
   '(2) reason-disagreements -> take one reviewer reason;',
   '(3) harmonized Non-health category for clearly non-health papers (engineering/IT/environmental), overriding "non-human" dropdowns;',
   '(4) STUDY-0097 folded from target-reached -> Non-human (screened at Seq 181; panel agreed Non-human);',
   '(5) STUDY-0683 recorded under YA label "narrative review" (factually a study protocol).',
   'STUDY-0219 (groundwater + health-risk) kept Non-human per user decision.']})

out='data/screening/08_13_2026_batch1_exclusions_LOCKED.xlsx'
with pd.ExcelWriter(out,engine='openpyxl') as xw:
    readme.to_excel(xw,sheet_name='README',index=False)
    summ.to_excel(xw,sheet_name='Tally',index=False)
    b1out.to_excel(xw,sheet_name='Batch1_audit_108',index=False)
    hf=Font(bold=True,color='FFFFFF'); fill=PatternFill('solid',fgColor='16697A')
    for sh in xw.sheets:
        ws=xw.sheets[sh]; ws.freeze_panes='A2'
        for c in ws[1]: c.font=hf; c.fill=fill; c.alignment=Alignment(vertical='top',wrap_text=True)
    xw.sheets['README'].column_dimensions['A'].width=26; xw.sheets['README'].column_dimensions['B'].width=118
    for c in xw.sheets['README']['A']: c.font=Font(bold=True); c.alignment=Alignment(vertical='top')
    for c in xw.sheets['README']['B']: c.alignment=Alignment(vertical='top',wrap_text=True)
    for shn,df in [('Tally',summ),('Batch1_audit_108',b1out)]:
        ws=xw.sheets[shn]
        for i,col in enumerate(df.columns,1):
            ws.column_dimensions[get_column_letter(i)].width=min(70,max(10,int(df[col].map(lambda v:len(str(v))).max())+2))
print("\nwrote",out)

# ---- update master CSV: batch-1 finalized, add STUDY-0097, mark status ----
rec['FINAL_reason']=rec.apply(lambda r: (B1[r['PMID']][0] if r['PMID'] in B1 else r['RESOLVED_reason']) if r['batch']=='batch-1' else r['RESOLVED_reason'], axis=1)
rec['status']=rec['batch'].map(lambda b:'batch-1 LOCKED' if b=='batch-1' else 'batch-2 provisional')
master=pd.concat([rec, extra.assign(status='batch-1 LOCKED')], ignore_index=True)
master.to_csv('data/screening/08_13_2026_exclusion_final_master.csv',index=False,encoding='utf-8-sig')
print("wrote data/screening/08_13_2026_exclusion_final_master.csv  (batch-1 LOCKED; batch-2 still provisional)")
