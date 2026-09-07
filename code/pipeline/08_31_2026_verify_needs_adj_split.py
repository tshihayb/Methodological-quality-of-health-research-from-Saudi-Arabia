# NOTE (public repository): 13 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
"""
VERIFY that the archived adjudication worklists follow from the raw reviewer exports.

Run:  python code/pipeline/08_31_2026_verify_needs_adj_split.py [--write-csv]
Exits 0 on pass, 1 on failure. Reads only private/reviewers/form-exports/ and the
archived worklists; writes nothing unless --write-csv is given.

WHY.  code/sas/05_19_2025_analysis.sas decided, for every paper and every item,
whether the two reviewers agreed. Everything downstream rests on that split, and
it has never been independently checked. The SAS cannot simply be rerun to prove
itself: its Needs_adj export was a blank worklist that TSA and YA then filled in
by hand, so the file of record is the annotated copy. What CAN be checked is the
decision rule, which is a pure function of the 14 exports.

SAS logic replicated (05_19_2025_analysis.sas:[ISSN-REDACTED]):
  1. stack r1..r14 in that order, each tagged Reviewer_ID = its number   (:2608, :132)
  2. delete Reviewer_ID = 3                                              (:2620)
  3. transpose by PMID; COL1 -> first_reviewer, COL2 -> second_reviewer,
     COL3 dropped                                                       (:2631, :2686)
  4. complete = 1 iff the Timestamp row has both COL1 and COL2 non-empty (:2696)
  5. Needs_adj      = complete=1 and first ne second                     (:2761)
     do_not_Need_adj = complete=1 and first  = second                    (:2897)
     both dropping four both-recused PMIDs                              (:2765, :2900)

STRICTLY READ-ONLY with respect to private/ and data/. Every write goes to
codes audit/output/rederive/. Verify with integrity_manifest.py.
"""
import openpyxl, os, re, csv, sys, collections, datetime

FE = 'private/reviewers/form-exports'
WL = 'private/reviewers/adjudication-worklists'
OUT = os.path.join('codes audit', 'output', 'rederive')   # only used with --write-csv

# Forms column order == SAS rename order (05_19_2025_analysis.sas:136-222),
# with Timestamp as form column 1.
CANON = [
    'Timestamp', 'PMID', 'recusal', 'task',
    'descriptive_design', 'descriptive_pop',
    'descriptive_sampling', 'descriptive_sample_size', 'descriptive_acc_sampl',
    'descriptive_sample_ach', 'descriptive_base_sel',
    'descriptive_outcome_type', 'descriptive_val_outcome', 'descriptive_out_bias_acc',
    'descriptive_miss_outcome', 'descriptive_hand_miss_outcom',
    'descriptive_err_disc', 'descriptive_confl_task',
    'predictive_design', 'predictive_pop',
    'causal_design', 'causal_pop',
    'causal_sampling', 'causal_sample_size', 'causal_acc_sampl',
    'causal_sample_ach', 'causal_base_sel', 'causal_comp_dis', 'causal_follow',
    'causal_ltfu_bias', 'causal_ltfu_acc',
    'causal_exposure_type', 'causal_val_exposure',
    'causal_diff_or_nondiff_exp', 'causal_exp_bias_acc',
    'causal_outcome_type', 'causal_val_outcome',
    'causal_diff_or_nondiff_out', 'causal_out_bias_acc',
    'causal_dep_or_indep_misc',
    'causal_base_conf_meth', 'causal_time_verying', 'causal_tv_conf_meth',
    'causal_conf_var_det',
    'causal_miss_exposure', 'causal_hand_miss_exposure',
    'causal_miss_outcome', 'causal_hand_miss_outcom',
    'causal_err_disc',
    'comments', 'comments_focus',
]

# PROC TRANSPOSE var list order (:[ISSN-REDACTED]). PMID is the BY key, not a row.
TVARS = [v for v in CANON if v not in ('Timestamp', 'PMID')] + ['Reviewer_ID', 'Timestamp']

DROP_REVIEWER = 3
DROP_PMIDS = {'STUDY-0956', 'STUDY-0275', 'STUDY-0320', 'STUDY-0927'}   # both reviewers recused

# Eight duplicate-submission resolutions, hand-written one per line INSIDE the
# individual reviewers' data steps rather than in one place, each keyed on PMID
# plus an exact timestamp string. Without these the transpose compares a reviewer
# against their own earlier submission. Line numbers are in 05_19_2025_analysis.sas.
DROP_ROWS = {
    ('STUDY-0183', '2025/02/02 7:13:59 PM GMT+3'),    # :403  r2 recusal, then resubmitted
    ('STUDY-0017', '2024/08/30 11:47:06 AM GMT+3'),   # :404  r2 duplicate
    ('STUDY-0171', '2025/01/14 11:00:41 AM GMT+3'),   # :760  r4 duplicate, 3 min apart
    ('STUDY-0449', '2025/05/16 10:55:46 PM GMT+3'),   # :1291 r7 duplicate
    ('STUDY-0852', '2024/12/07 9:06:59 PM GMT+3'),    # :1476 r8 duplicate
    ('STUDY-0275', '2024/12/16 8:58:20 PM GMT+3'),    # :1477 r8, same data step as above
    ('STUDY-0228', '2024/12/22 11:58:30 AM GMT+3'),   # :1660 r9
    ('STUDY-0336', '2024/08/15 10:42:46 AM GMT+3'),   # :2018 r11
}

# One PMID repair, also hand-written and also buried in a reviewer's data step (:1292):
#     if What_was_the_PMID_of_the_researc=. then What_was_the_PMID_of_the_researc=STUDY-0473;
# Reviewer 7 typed the task answer, "Descriptive", into the PMID field of the form.
# PROC IMPORT reads that column as numeric, so the text arrived as a missing value, and
# the analyst filled it in from context. The patch is keyed on missingness, not on the
# row, so it is safe only while that reviewer has exactly one unparseable PMID. Across
# all 14 exports there is exactly one, so it resolves correctly here.
PMID_REPAIR = {7: 'STUDY-0473'}


def norm(v):
    """SAS character semantics: missing is '', trailing blanks are not significant."""
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return str(v).rstrip()


def pmid_of(v):
    """SAS converts the PMID column to numeric (:43-45), which strips surrounding
    whitespace and turns non-numeric text into missing. One raw cell carries a
    leading tab and one holds the word 'Descriptive'; treating those as literal
    keys split one paper in two and invented another."""
    s = norm(v).strip()
    m = re.match(r'^(\d+)', s.replace('.0', ''))
    return m.group(1) if m else ''


def load_exports():
    files = {}
    for fn in os.listdir(FE):
        m = re.search(r'Reviewer_(\d+)\.xlsx$', fn, re.I)
        if m and not fn.startswith('~$'):
            files[int(m.group(1))] = os.path.join(FE, fn)
    assert len(files) == 14, 'expected 14 exports, found %d' % len(files)

    headers, stacked, dropped_rows, repaired = {}, [], [], []
    for rid in sorted(files):                       # r1..r14, the SAS stack order
        wb = openpyxl.load_workbook(files[rid], read_only=True, data_only=True)
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        hdr = [norm(c) for c in next(it)]
        headers[rid] = hdr
        for row in it:
            if row is None or len(row) < 2:
                continue
            rec = {CANON[i]: row[i] for i in range(min(len(CANON), len(row)))}
            rec['Reviewer_ID'] = rid
            # The repair must run BEFORE the unparseable-PMID guard below, or the very
            # row it exists to rescue is discarded on the way in.
            if (pmid_of(rec['PMID']) == '' and norm(rec['PMID']) != ''
                    and rid in PMID_REPAIR):
                repaired.append((rid, norm(rec['PMID']), PMID_REPAIR[rid]))
                rec['PMID'] = PMID_REPAIR[rid]
            if pmid_of(rec['PMID']) == '':
                continue
            if (pmid_of(rec['PMID']), norm(rec.get('Timestamp'))) in DROP_ROWS:
                dropped_rows.append((rid, pmid_of(rec['PMID']), norm(rec.get('Timestamp'))))
                continue
            stacked.append(rec)
        wb.close()

    ref = headers[1]
    bad = [r for r in headers if headers[r] != ref]
    print('header check: %d/%d exports share reviewer 1\'s %d-column header%s'
          % (14 - len(bad), 14, len(ref), '' if not bad else '  MISMATCH: %s' % bad))
    if bad:
        for r in bad:
            for i, (a, b) in enumerate(zip(ref, headers[r])):
                if a != b:
                    print('   r%-2d col %2d: %r vs %r' % (r, i + 1, a[:40], b[:40]))
    assert len(ref) == 51, 'expected 51 columns, got %d' % len(ref)
    for rid, was, now in repaired:
        print('PMID repair applied: r%s  %r -> %s' % (rid, was, now))
    print('duplicate-submission deletions applied: %d of %d'
          % (len(dropped_rows), len(DROP_ROWS)))
    for rid, p, ts in dropped_rows:
        print('   r%-3s %s  %s' % (rid, p, ts))
    missed = DROP_ROWS - {(p, ts) for _, p, ts in dropped_rows}
    for p, ts in sorted(missed):
        print('   !! NOT MATCHED in the raw exports: %s  %s' % (p, ts))
    return stacked


def transpose(stacked):
    """SAS PROC TRANSPOSE BY PMID: COL1/COL2 from stack order, COL3 discarded."""
    by = collections.OrderedDict()
    for rec in stacked:
        by.setdefault(pmid_of(rec['PMID']), []).append(rec)

    rows, dropped3 = [], 0
    for pmid, recs in by.items():
        kept = [r for r in recs if r['Reviewer_ID'] != DROP_REVIEWER]
        dropped3 += len(recs) - len(kept)
        first = kept[0] if len(kept) > 0 else None
        second = kept[1] if len(kept) > 1 else None
        ts1 = norm(first.get('Timestamp')) if first else ''
        ts2 = norm(second.get('Timestamp')) if second else ''
        complete = 1 if (ts1 != '' and ts2 != '') else 0
        for v in TVARS:
            rows.append({
                'PMID': pmid, 'variable': v,
                'first_reviewer': norm(first.get(v)) if first else '',
                'second_reviewer': norm(second.get(v)) if second else '',
                'complete': complete,
            })
    print('reviewer %d rows dropped: %d | papers: %d' % (DROP_REVIEWER, dropped3, len(by)))
    return rows


def split(rows):
    needs, dont = [], []
    for r in rows:
        if r['complete'] != 1 or r['PMID'] in DROP_PMIDS:
            continue
        (needs if r['first_reviewer'] != r['second_reviewer'] else dont).append(r)
    return needs, dont


def load_sas(path, keycols=('PMID', 'variable')):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    hdr = [norm(c) for c in next(it)]
    ci = {c: i for i, c in enumerate(hdr)}
    out = {}
    for row in it:
        if row is None or ci.get('PMID') is None or row[ci['PMID']] is None:
            continue
        k = (pmid_of(row[ci['PMID']]), norm(row[ci['variable']]))
        out[k] = (norm(row[ci['first_reviewer']]) if 'first_reviewer' in ci else '',
                  norm(row[ci['second_reviewer']]) if 'second_reviewer' in ci else '')
    wb.close()
    return out


def diff(label, mine, theirs, show=12):
    mk, tk = set(mine), set(theirs)
    only_mine, only_theirs = mk - tk, tk - mk
    shared = mk & tk
    valdiff = [k for k in shared if mine[k] != theirs[k]]
    print('\n---- %s ----' % label)
    print('  re-derived rows : %d' % len(mk))
    print('  SAS rows        : %d' % len(tk))
    print('  in both         : %d' % len(shared))
    print('  ONLY re-derived : %d' % len(only_mine))
    print('  ONLY SAS        : %d' % len(only_theirs))
    print('  value mismatches: %d' % len(valdiff))
    for k in sorted(only_mine)[:show]:
        print('     only mine  %s / %s' % k)
    for k in sorted(only_theirs)[:show]:
        print('     only SAS   %s / %s' % k)
    for k in sorted(valdiff)[:show]:
        print('     differs    %s / %s' % k)
        print('        mine %r | %r' % mine[k])
        print('        SAS  %r | %r' % theirs[k])
    ok = not only_mine and not only_theirs and not valdiff
    print('  %s' % ('MATCH' if ok else 'DIFFERENCES PRESENT'))
    return ok


def write_csv(path, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=['PMID', 'variable', 'first_reviewer',
                                          'second_reviewer', 'complete'])
        w.writeheader()
        w.writerows(rows)


def classify(mine, theirs):
    """Split the value differences into the two known-benign classes and the rest.

    BENIGN 1  the worklist value is a strict prefix cut at 254/255 characters, which is
              PROC IMPORT DBMS=EXCEL capping character columns during the carry-forward
              re-import (:2840). The full text survives in the raw exports.
    BENIGN 2  paper STUDY-0031 carries TWO Timestamp rows in the worklist, one naming the
              reviewer who withdrew and one the reviewer who replaced him, a leftover of
              the replacement. Which one is read depends on row order, so this compares
              unequal without either being wrong. Housekeeping only: Timestamp is dropped
              before scoring (:2934) and before the phase II comparison.
    Anything else is a real difference and fails the check.
    """
    trunc, stale, real = 0, 0, []
    for k in set(mine) & set(theirs):
        if mine[k] == theirs[k]:
            continue
        for a, b in zip(mine[k], theirs[k]):
            if a == b:
                continue
            if k[1] == 'comments' and a.startswith(b) and len(b) in (254, 255):
                trunc += 1
            elif k == ('STUDY-0031', 'Timestamp'):
                stale += 1
            else:
                real.append((k, a, b))
    return trunc, stale, real


def main():
    write = '--write-csv' in sys.argv
    print('=' * 78)
    print('VERIFY: does the needs-adjudication split follow from the raw reviewer exports?')
    print('=' * 78)
    stacked = load_exports()
    print('stacked rows (all 14): %d' % len(stacked))
    needs, dont = split(transpose(stacked))
    print('\nre-derived: Needs_adj %d rows | do_not_Need_adj %d rows' % (len(needs), len(dont)))

    if write:
        os.makedirs(OUT, exist_ok=True)
        write_csv(os.path.join(OUT, 'rederived_Needs_adj.csv'), needs)
        write_csv(os.path.join(OUT, 'rederived_do_not_Need_adj.csv'), dont)
        print('csvs -> %s' % OUT.replace(os.sep, '/'))

    mine_d = {(r['PMID'], r['variable']): (r['first_reviewer'], r['second_reviewer'])
              for r in dont}
    sas_d = load_sas(os.path.join(WL, '04_17_2026_do_not_Need_adj.xlsx'))
    ok1 = diff('do_not_Need_adj  vs  04_17_2026_do_not_Need_adj.xlsx', mine_d, sas_d)

    mine_n = {(r['PMID'], r['variable']): (r['first_reviewer'], r['second_reviewer'])
              for r in needs}
    sas_n = load_sas(os.path.join(WL, '04_17_2026_Needs_adj_raw_progress_TA.xlsx'))
    keys_ok = set(mine_n) == set(sas_n)
    trunc, stale, real = classify(mine_n, sas_n)

    print('\n---- Needs_adj  vs  04_17_2026_Needs_adj_raw_progress_TA.xlsx ----')
    print('  rows %d vs %d | same row set: %s' % (len(mine_n), len(sas_n), keys_ok))
    print('  known-benign 255-char comment truncations : %d' % trunc)
    print('  known-benign stale Timestamp (STUDY-0031)   : %d' % stale)
    print('  UNEXPLAINED differences                   : %d' % len(real))
    for (k, a, b) in real[:10]:
        print('     %s / %s\n        raw      %r\n        worklist %r' % (k[0], k[1], a, b))

    passed = ok1 and keys_ok and not real
    print('\n' + '=' * 78)
    if passed:
        print('PASS: the split follows from the raw exports, row for row and value for value.')
        print('      The only differences are %d comment truncations and 1 stale timestamp,' % trunc)
        print('      none of which is a scored item (dropped at 05_19_2025_analysis.sas:2934).')
    else:
        print('FAIL: the archived worklists no longer follow from the raw exports.')
        print('      Do not rebuild downstream artifacts until this is understood.')
    return 0 if passed else 1


if __name__ == '__main__':
    sys.exit(main())
