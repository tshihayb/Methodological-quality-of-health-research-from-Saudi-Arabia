"""
Re-sort private/reviewers/adjudication-worklists/04_17_2026_Needs_adj_raw_progress_TA.xlsx so that, within each PMID,
rows follow the data-collection tool's order of appearance (Sections 1-19)
instead of alphabetical order of `variable`.

Canonical order source (two independent, mutually confirming sources):
  1. code/sas/05_19_2025_analysis.sas rename block, explicitly grouped by tool Section 1-19
  2. The 51-column header shared identically by all 14 Google-Forms reviewer exports
     (Forms exports preserve question order as column order)

Pure reorder: values and per-cell styling are carried with their rows, nothing edited.
"""
import openpyxl, os
from copy import copy

P = r"."
SRC = os.path.join(P, "private/reviewers/adjudication-worklists/04_17_2026_Needs_adj_raw_progress_TA.xlsx")
DST = os.path.join(P, "private/reviewers/adjudication-worklists/07_19_2026_Needs_adj_progress_TA_toolorder.xlsx")

# ---- canonical tool order: (section, variable) -------------------------------
# PMID + recusal (Section 1) are not rows in this file: PMID is the key column and
# recusal never produced a disagreement. Reviewer_ID is derived metadata (SAS sets it
# in the data step); it is kept with Timestamp as the leading housekeeping pair.
TOOL_ORDER = [
    (0,  "Reviewer_ID"),                    # derived metadata
    (0,  "Timestamp"),                      # form column 1
    (2,  "task"),                           # Section 2
    (3,  "descriptive_design"),             # Section 3
    (3,  "descriptive_pop"),
    (4,  "descriptive_sampling"),           # Section 4
    (4,  "descriptive_sample_size"),
    (4,  "descriptive_acc_sampl"),
    (4,  "descriptive_sample_ach"),
    (4,  "descriptive_base_sel"),
    (5,  "descriptive_outcome_type"),       # Section 5
    (5,  "descriptive_val_outcome"),
    (5,  "descriptive_out_bias_acc"),
    (6,  "descriptive_miss_outcome"),       # Section 6
    (6,  "descriptive_hand_miss_outcom"),
    (7,  "descriptive_err_disc"),           # Section 7
    (8,  "descriptive_confl_task"),         # Section 8
    (9,  "predictive_design"),              # Section 9
    (9,  "predictive_pop"),
    (10, "causal_design"),                  # Section 10
    (10, "causal_pop"),
    (11, "causal_sampling"),                # Section 11
    (11, "causal_sample_size"),
    (11, "causal_acc_sampl"),
    (11, "causal_sample_ach"),
    (11, "causal_base_sel"),
    (11, "causal_comp_dis"),
    (11, "causal_follow"),
    (11, "causal_ltfu_bias"),
    (11, "causal_ltfu_acc"),
    (12, "causal_exposure_type"),           # Section 12
    (12, "causal_val_exposure"),
    (12, "causal_diff_or_nondiff_exp"),
    (12, "causal_exp_bias_acc"),
    (13, "causal_outcome_type"),            # Section 13
    (13, "causal_val_outcome"),
    (13, "causal_diff_or_nondiff_out"),
    (13, "causal_out_bias_acc"),
    (14, "causal_dep_or_indep_misc"),       # Section 14
    (15, "causal_base_conf_meth"),          # Section 15
    (15, "causal_time_verying"),
    (15, "causal_tv_conf_meth"),
    (15, "causal_conf_var_det"),
    (16, "causal_miss_exposure"),           # Section 16
    (16, "causal_hand_miss_exposure"),
    (17, "causal_miss_outcome"),            # Section 17
    (17, "causal_hand_miss_outcom"),
    (18, "causal_err_disc"),                # Section 18
    (19, "comments"),                       # Section 19
    (19, "comments_focus"),
]
RANK = {v: i for i, (s, v) in enumerate(TOOL_ORDER)}

# ---- load ---------------------------------------------------------------------
wb = openpyxl.load_workbook(SRC)
ws = wb.worksheets[0]
hdr = [c.value for c in ws[1]]
iP, iV = hdr.index("PMID"), hdr.index("variable")

body = list(ws.iter_rows(min_row=2))

unknown = sorted({r[iV].value for r in body} - set(RANK))
if unknown:
    raise SystemExit(f"ABORT - variables absent from the tool-order map: {unknown}")

# PMID blocks keep their existing first-appearance sequence (no unrequested resort)
pmid_seq, seen = [], set()
for r in body:
    p = r[iP].value
    if p not in seen:
        seen.add(p); pmid_seq.append(p)
PPOS = {p: i for i, p in enumerate(pmid_seq)}

# stable sort: PMID block, then tool order, then original position for any ties
order = sorted(range(len(body)),
               key=lambda i: (PPOS[body[i][iP].value], RANK[body[i][iV].value], i))

# ---- write --------------------------------------------------------------------
# Permute rows *inside the loaded workbook*: a cell's _style is an index into that
# workbook's own fill/font/border tables, so it is only valid in this workbook.
# Snapshot every row first, then lay the snapshots back down in the new order.
snap = [[(c.value, copy(c._style), c.hyperlink) for c in r] for r in body]

for newrow, i in enumerate(order, start=2):
    for c, (val, style, link) in zip(ws[newrow], snap[i]):
        c.value = val
        c._style = style
        c.hyperlink = copy(link) if link else None

ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(ws.max_column)}{len(body)+1}"
ws.freeze_panes = "A2"                            # header stays visible while scrolling

wb.save(DST)
print(f"wrote {DST}\n  {len(body)} body rows | {len(pmid_seq)} PMIDs")
