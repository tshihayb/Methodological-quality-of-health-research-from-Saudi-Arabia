# -*- coding: utf-8 -*-
# Append the 39 batch-6 phase-I disagreement rows to copies of the existing TA and YA
# needs-adjudication progress worklists (preserving their completed adjudications).
import csv, warnings
warnings.simplefilter("ignore")
from openpyxl import load_workbook, Workbook

new=[c for c in csv.DictReader(open("data/adjudication/08_08_2026_batch6_cell_detail.csv",encoding="utf-8-sig")) if c["needs_adj"]=="True"]
# metadata for the 8 papers (Title/exposure/outcome) from the batch6 Needs_adj we already built
wb=load_workbook("private/reviewers/adjudication-worklists/08_08_2026_batch6_Needs_adj.xlsx"); ws=wb.active
b6=list(ws.iter_rows(values_only=True)); b6h=list(b6[0]); wb.close()
metarow={}
for r in b6[1:]:
    d=dict(zip(b6h,r)); metarow[(str(d["PMID"]),d["variable"])]=d

def append_to(src, out, who):
    wb=load_workbook(src); ws=wb.active
    hdr=[c.value for c in ws[1]]
    ci={h:i for i,h in enumerate(hdr)}
    added=0
    for c in new:
        d=metarow[(c["PMID"],c["variable"])]
        row=[""]*len(hdr)
        for k in ["PMID","Title","Study_Type","adjudicated_exposure","adjudicated_outcome","variable","first_reviewer","second_reviewer","complete"]:
            if k in ci: row[ci[k]]=d[k]
        ws.append(row); added+=1
    wb.save(out)
    print(f"{out}: appended {added} rows -> total data rows {ws.max_row-1}")

append_to("private/reviewers/adjudication-worklists/07_22_2026_Needs_adj_progress_TA.xlsx","08_08_2026_Needs_adj_progress_TA.xlsx","TA")
append_to("private/reviewers/adjudication-worklists/07_22_2026_Needs_adj_progress_YA.xlsx","08_08_2026_Needs_adj_progress_YA.xlsx","YA")
