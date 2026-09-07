# NOTE (public repository): 32 hard-coded study identifier(s) were
# rewritten to de-identified surrogates by tools/build_public_repo.py.
# They are rulings about particular studies rather than values the data
# can supply, so they have to travel with it.
# -*- coding: utf-8 -*-
# Consolidated inconsistency RESOLUTION WORKLIST for TSA + YA, on the final 385 dataset.
# Every pending item: current value, plain issue, researched proposal, blank decision column.
import csv, importlib.util
import warnings; warnings.simplefilter("ignore")
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

LONG="data/analysis/08_08_2026_ANALYSIS_DATASET_long_385.csv"

# ---- run checker on 385 to get the live GAPs + soft violations ----
spec=importlib.util.spec_from_file_location("chk","code/pipeline/_check_contradictions.py")
chk=importlib.util.module_from_spec(spec); spec.loader.exec_module(chk)
W=chk.load(LONG); R=chk.soft(W); gaps,extras=chk.skiplogic(W)

# ---- current values lookup (pivoted) ----
val={}
for r in csv.DictReader(open(LONG,encoding="utf-8-sig")):
    val[(r["PMID"],r["variable"])]=(r["final"] or "").strip()
def gv(pmid,short,task="causal"):
    return val.get((pmid,f"{task}_{short}"),"")

# ---- researched resolutions (from the 2026-07-22/23 full-text screens + rulings) ----
# follow-up
FOLLOW={
 "STUDY-0093":("Yes","No","ecological / hospital units, no persons followed"),
 "STUDY-0331":("Yes","No","tx-stage groups -> Candida = different patients"),
 "STUDY-0202":("Yes","No","CCT->IOP is a cross-sectional correlation; the TPRK before/after is a different pair"),
 "STUDY-0326":("No","Yes","Age->Hyponatremia, incident outcome tracked daily"),
 "STUDY-0471":("No","Yes","Music->anxiety, same 18 people pre/post"),
}
DESIGN_XS={"STUDY-0845","STUDY-0306","STUDY-0878","STUDY-0676"}   # S22 Cohort->Cross-sectional
TORCT={"STUDY-0431","STUDY-0385"}                              # S6a Cohort->RCT
S5_GENUINE={"STUDY-0597","STUDY-0429","STUDY-0923","STUDY-0736","STUDY-0695","STUDY-0919"}

rows=[]  # (category, rule, pmid, item, current, issue, proposal, status, notes)
def add(cat,rule,pmid,item,cur,issue,prop,status,notes=""):
    rows.append([cat,rule,pmid,item,cur,issue,prop,status,"",notes])

# ---------- 1. HARD skip-logic GAPs ----------
for t,pmid,parent,pval,child,cval in gaps:
    short=child.split("_",1)[1] if "_" in child else child
    if "hand_miss" in child:
        prop="No"; status="proposed - verify"; issue=f"{parent}={pval[:22]} makes '{child}' apply, but it is blank (no handling recorded)"
        notes="missing-data handling: 'No' (no imputation) in 19/24 & 14/17 of comparable papers"
    else:
        prop=""; status="NEEDS JUDGEMENT (full text)"; issue=f"{parent}={pval[:22]} makes '{child}' apply, but it is blank"
        notes={"acc_sampl":"did they account for non-response in the sample-size calc?","sample_ach":"was the required sample size achieved?",
               "ltfu_acc":"was the loss-to-follow-up bias accounted for?","comp_dis":"did they compare responders vs non-responders?"}.get(short,"")
    add("1. Skip-logic GAP (fill the blank)","GAP",pmid,child,"(blank)",issue,prop,status,notes)

# ---------- 2. Follow-up contradictions (S1/S8/S22) ----------
s1=set(R.get("S1",[]))|set(R.get("S8",[])); s22=set(R.get("S22",[]))
for pmid in sorted(s1|s22):
    if pmid in FOLLOW:
        cur_f,new_f,why=FOLLOW[pmid]
        add("2. Follow-up (design vs follow-up)","S1/S8" if pmid in s1 else "S22",pmid,"follow",
            f"follow={gv(pmid,'follow')}, design={gv(pmid,'design')}",
            "design and follow-up are logically inconsistent",
            f"flip follow {cur_f} -> {new_f}","RESEARCHED - ratify",why)
    elif pmid in DESIGN_XS:
        add("2. Follow-up (design vs follow-up)","S22",pmid,"design",
            f"design={gv(pmid,'design')}, follow={gv(pmid,'follow')}",
            "longitudinal design with follow=No",
            "relabel design Cohort -> Cross-sectional","RESEARCHED - ratify","selected pair is cross-sectional (full-text)")
    elif pmid=="STUDY-0792":
        add("2. Follow-up (design vs follow-up)","S22",pmid,"follow",
            f"design={gv(pmid,'design')}, follow={gv(pmid,'follow')}","RCT with follow=No",
            "NO CHANGE","confirm - no change","immediate one-time outcome (VR->venipuncture pain); follow=No correct")
    else:  # STUDY-0475, STUDY-0679 = new
        add("2. Follow-up (design vs follow-up)","S22",pmid,"design/follow",
            f"design={gv(pmid,'design')}, follow={gv(pmid,'follow')}","longitudinal design with follow=No (NEW paper)",
            "(pending Claude full-text screen)","NEEDS SCREENING (Claude)","not yet screened")

# ---------- 3. Design / randomization (S6a, S6b) ----------
for pmid in sorted(set(R.get("S6a",[]))):
    add("3. Design & randomization","S6a",pmid,"design",f"design={gv(pmid,'design')}, base_conf_meth={gv(pmid,'base_conf_meth')[:26]}",
        "'Randomization' used as a method but design is not RCT","relabel design Cohort -> RCT","RESEARCHED - ratify",
        "confirmed randomized for the selected exposure (full-text)")
for pmid in sorted(set(R.get("S6b",[]))):
    if pmid=="STUDY-0499":
        add("3. Design & randomization","S6b",pmid,"base_conf_meth",f"design={gv(pmid,'design')}, base_conf_meth={gv(pmid,'base_conf_meth')[:26]}",
            "RCT but 'Randomization' not listed among methods","NO CHANGE","confirm - no change",
            "double-blind RCT; 'No adjustment' = randomization-only convention (Ruling 1)")
    else:
        add("3. Design & randomization","S6b",pmid,"base_conf_meth",f"design={gv(pmid,'design')}, base_conf_meth={gv(pmid,'base_conf_meth')[:26]}",
            "RCT recording 'No adjustment' (NEW paper)","likely randomization-only convention; confirm via ITT/PP screen","NEEDS SCREENING (Claude)","")

# ---------- 4. Sampling vs selection bias (S10/S11) ----------
for tag in ["C","D"]:
    for pmid in sorted(set(R.get("S10"+tag,[]))):
        task="descriptive" if tag=="D" else "causal"
        cur=f"sampling={gv(pmid,'sampling',task)}, base_sel={gv(pmid,'base_sel',task)[:30]}, design={gv(pmid,'design',task)}"
        if pmid=="STUDY-0520":
            add("4. Sampling vs selection bias","S10/S11",pmid,"base_sel",cur,
                "non-random sample called 'No baseline selection bias'","APPLY: change base_sel -> selection bias PRESENT",
                "GENUINE ERROR - apply","HSV seroprevalence from a clinically-tested (unrepresentative) sample")
        elif pmid=="STUDY-0084":
            add("4. Sampling vs selection bias","S10/S11",pmid,"base_sel",cur,
                "non-random sample + 'No baseline selection bias' (NEW paper)","likely RCT false alarm; confirm","NEEDS SCREENING (Claude)","DELIVER trial")
        else:
            note = "RCT: randomization handles internal validity" if pmid!="STUDY-0345" else "30-yr tumour-registry census (complete)"
            if pmid=="STUDY-0700": note="base_sel already normalised in phase II"
            add("4. Sampling vs selection bias","S10/S11",pmid,"base_sel",cur,
                "non-random sample + 'No baseline selection bias'","NO CHANGE (false alarm)","confirm - no change",note)

# ---------- 5. Confounding basis (S5) ----------
for pmid in sorted(set(R.get("S5",[]))):
    cur=f"base_conf_meth={gv(pmid,'base_conf_meth')[:26]}, conf_var_det={gv(pmid,'conf_var_det')[:26]}"
    if pmid in S5_GENUINE:
        add("5. Confounding basis","S5",pmid,"conf_var_det",cur,
            "a real confounding method was used but no basis for choosing confounders is named",
            "name a confounder-selection basis, or confirm it as a real error","GENUINE - decide","")
    elif pmid=="STUDY-0613":
        add("5. Confounding basis","S5",pmid,"conf_var_det",cur,"real method, no basis named",
            "leave as-is (you agreed it is unknowable)","confirm - no change","the one phase-II cell you deliberately left blank")
    else:
        add("5. Confounding basis","S5",pmid,"conf_var_det",cur,"real method, no basis named",
            "N/A - randomization-only trial (moot under Ruling 3)","confirm - no change","conf_var_det is N/A for randomization-only trials")

# ---------- 6. Rulings TSA (+YA) owe ----------
for rule,item,issue,prop in [
 ("Ruling","val_exposure","Is exposure validity N/A for randomized interventions? (Finding 2, recommended, unruled)","Rule N/A for RCTs (regularises 9 papers, removes 3 measurement flags)"),
 ("Ruling","conf_var_det","Manuscript framing: report conf_var_det N/A-share vs exemption-share","Pick one framing for the manuscript"),
 ("Ruling","rct_pp_unadjusted","Where does the per-protocol-unadjusted confounding error live? (STUDY-0532, STUDY-0062)","Decide the scoring home"),
 ("Ruling","analysis_population","How to score 'Not reported' analysis population","Decide the scoring rule"),
]:
    add("6. Rulings (project-level)",rule,"(all / noted)",item,"-",issue,prop,"TSA + YA DECISION","")

# ---------- 7. Extraction Claude runs, TSA/YA verify ----------
add("7. Extraction (Claude runs)","ITT/PP","STUDY-0431, STUDY-0385, STUDY-1022, STUDY-0159, STUDY-0963, STUDY-0834, STUDY-0084",
    "rct_analysis_pop","-","7 RCTs need the ITT vs per-protocol analysis-population screen from full text",
    "Claude extracts; TSA/YA verify","CLAUDE TO RUN","feeds Ruling 1 (ITT indicator)")

# ---------- write Excel ----------
wb=Workbook(); ws=wb.active; ws.title="resolution_worklist"
HDR=["#","Category","Rule","PMID","Item","Current value","Issue","Proposed resolution (researched)","Status","TSA + YA decision","Notes"]
ws.append(HDR)
for i,r in enumerate(rows,1): ws.append([i]+r)

# styling
hfill=PatternFill("solid",fgColor="16697A"); hfont=Font(bold=True,color="FFFFFF",size=10)
decfill=PatternFill("solid",fgColor="FFF3D6"); thin=Side(style="thin",color="D9D9D9")
border=Border(left=thin,right=thin,top=thin,bottom=thin)
catcol={"1":"E8F1F2","2":"FCEDE4","3":"EDE7F3","4":"E7F0EA","5":"F3ECDA","6":"F5E1E1","7":"E4EEF0"}
for c in range(1,len(HDR)+1):
    cell=ws.cell(1,c); cell.fill=hfill; cell.font=hfont; cell.alignment=Alignment(vertical="center",wrap_text=True); cell.border=border
widths=[4,30,8,22,22,34,46,44,22,26,40]
for c,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(c)].width=w
statuscolor={"GENUINE - decide":"C00000","GENUINE ERROR - apply":"C00000","NEEDS JUDGEMENT (full text)":"BF8F00",
             "TSA + YA DECISION":"C55A11","NEEDS SCREENING (Claude)":"7030A0","CLAUDE TO RUN":"7030A0",
             "RESEARCHED - ratify":"1F7A1F","proposed - verify":"1F7A1F","confirm - no change":"808080"}
for ri in range(2,ws.max_row+1):
    cat=ws.cell(ri,2).value; key=cat.split(".")[0]
    fill=PatternFill("solid",fgColor=catcol.get(key,"FFFFFF"))
    for c in range(1,len(HDR)+1):
        cell=ws.cell(ri,c); cell.border=border; cell.alignment=Alignment(vertical="top",wrap_text=True)
        if c!=10: cell.fill=fill
    ws.cell(ri,10).fill=decfill                       # decision column highlighted
    st=ws.cell(ri,9); st.font=Font(bold=True,color=statuscolor.get(st.value,"000000"),size=9)
ws.freeze_panes="A2"; ws.auto_filter.ref=f"A1:{get_column_letter(len(HDR))}{ws.max_row}"

# summary sheet
sm=wb.create_sheet("summary")
from collections import Counter
cats=Counter(r[0] for r in rows); stats=Counter(r[7] for r in rows)
sm.append(["INCONSISTENCY RESOLUTION WORKLIST - 385 papers (377 + 8 batch-6)"])
sm.append([f"Total items to work through: {len(rows)}"]); sm.append([])
sm.append(["By category:"])
for k in sorted(cats): sm.append(["",k,cats[k]])
sm.append([]); sm.append(["By status / who acts:"])
for k,v in stats.most_common(): sm.append(["",k,v])
sm.column_dimensions["B"].width=42
sm["A1"].font=Font(bold=True,size=12)

wb.save("data/quality-control/08_08_2026_inconsistency_resolution_worklist.xlsx")
print(f"wrote data/quality-control/08_08_2026_inconsistency_resolution_worklist.xlsx  ({len(rows)} items)")
print("By category:")
for k in sorted(cats): print(f"  {cats[k]:>3}  {k}")
print("By status:")
for k,v in stats.most_common(): print(f"  {v:>3}  {k}")
